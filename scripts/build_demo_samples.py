"""生成一次性虚构样本与待审核标注；已有目录绝不覆盖。"""
from pathlib import Path
import hashlib
import json

ROOT = Path(__file__).resolve().parents[1] / "tests/fixtures/demo"
PROFILES = [
    ("林舟", "Python、FastAPI、SQL、pytest", "项目：独立编写 FastAPI 订单接口，SQL 事务防止重复下单，pytest 覆盖异常回滚。", "实习：虚构甲公司后端开发三个月。", [80,95], [15,40]),
    ("陈澄", "TypeScript、React、CSS、Playwright", "项目：独立完成 React 管理后台，适配手机、键盘操作并用 Playwright 回归。", "实习：虚构乙公司前端开发三个月。", [15,40], [80,95]),
    ("周禾", "Python、FastAPI、SQL、TypeScript、React", "项目：完成 FastAPI 与 React 课程系统，负责接口和页面，没有压测或自动化测试记录。", "经历：课程项目，无企业实习。", [60,80], [60,80]),
    ("许青", "Python、SQL", "项目：用 Python 清洗 CSV 并用 SQL 汇总数据；没有 Web API 项目。", "经历：校内数据分析课程。", [30,55], [5,25]),
    ("沈夏", "React、JavaScript、CSS", "项目：React 静态作品集，支持移动端，未说明键盘导航和测试。", "经历：个人练习，无企业实习。", [5,25], [45,70]),
    ("顾言", "Java、Spring Boot、SQL、JUnit", "项目：用 Spring Boot 编写库存接口、事务与 JUnit 测试；未使用 FastAPI。", "实习：虚构丙公司 Java 后端两个月。", [40,65], [5,30]),
    ("叶初", "Python、PyTorch", "项目：训练图片分类模型，记录验证集精度；未部署 Web 服务。", "经历：实验室课程项目。", [20,45], [5,25]),
    ("何悠", "Figma、视觉设计", "项目：制作活动海报与界面设计稿，没有软件开发经历。", "经历：设计社团。", [0,20], [10,35]),
    ("唐星", "Python、FastAPI、SQL、pytest", "项目：为虚构商城编写 API 与回归用例，说明事务回滚边界。", "经历：个人项目，未说明可实习时间。", [65,85], [10,30]),
    ("宋溪", "TypeScript、React、CSS", "项目：React 表单校验和移动布局，提供键盘导航说明，未说明测试。", "经历：校园项目，无企业实习。", [10,30], [60,80]),
    ("陆远", "Python、FastAPI、SQL", "项目：两人合作实现任务 API；仅负责查询接口，不能把队友的鉴权工作算作本人实现。", "经历：团队课程项目，无测试证据。", [40,65], [10,30]),
    ("江月", "TypeScript、React、CSS、Playwright", "项目：在已有 React 项目修复表单无障碍缺陷并新增 Playwright 回归，保留变更记录。", "实习：虚构丁公司前端两个月。", [10,35], [75,95]),
]
JOBS = [
    {"id":"backend", "title":"虚构后端实习生", "requirements":["Python 与 FastAPI 接口开发", "SQL 与事务", "接口测试和异常边界", "说明本人负责的实现", "每周至少四天实习"]},
    {"id":"frontend", "title":"虚构前端实习生", "requirements":["TypeScript 与 React", "CSS 与移动端适配", "键盘无障碍和浏览器测试", "说明本人负责的实现", "每周至少四天实习"]},
]


def text_pdf(path, lines):
    """无需额外字体下载的两页文本 PDF，供提取器回归使用。"""
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, ArrayObject, TextStringObject, NumberObject, DecodedStreamObject
    n=NameObject
    writer=PdfWriter()
    descendant=DictionaryObject({n('/Type'):n('/Font'),n('/Subtype'):n('/CIDFontType0'),n('/BaseFont'):n('/STSong-Light'),
        n('/CIDSystemInfo'):DictionaryObject({n('/Registry'):TextStringObject('Adobe'),n('/Ordering'):TextStringObject('GB1'),n('/Supplement'):NumberObject(4)})})
    font=writer._add_object(DictionaryObject({n('/Type'):n('/Font'),n('/Subtype'):n('/Type0'),n('/BaseFont'):n('/STSong-Light'),
        n('/Encoding'):n('/UniGB-UCS2-H'),n('/DescendantFonts'):ArrayObject([writer._add_object(descendant)])}))
    for chunk in (lines[:5],lines[5:]):
        page=writer.add_blank_page(width=595,height=842)
        page[n('/Resources')]=DictionaryObject({n('/Font'):DictionaryObject({n('/F1'):font})})
        content='BT /F1 9 Tf 42 790 Td 24 TL\n'+'\n'.join('<'+line.encode('utf-16-be').hex()+'> Tj T*' for line in chunk)+'\nET'
        stream=DecodedStreamObject(); stream.set_data(content.encode('ascii'))
        page[n('/Contents')]=writer._add_object(stream)
    writer.write(path)


def build():
    ROOT.mkdir(parents=True, exist_ok=False)
    samples, labels = [], []
    for job in JOBS:
        (ROOT / f"jd-{job['id']}.txt").write_text(job["title"]+"\n要求："+"；".join(job["requirements"]), encoding="utf-8")
    for index, (name, skills, project, experience, backend, frontend) in enumerate(PROFILES, 1):
        sample_id = f"r{index:02d}"
        identity = {"name":name,"phone":f"1390000{index:04d}","email":f"synthetic-{sample_id}@example.test"}
        # 09 缺邮箱、10 缺手机、11 缺姓名：由明确的虚构人工补充资料修复。
        missing_field = {9:"email",10:"phone",11:"name"}.get(index)
        lines = [f"{label}：{identity[key]}" for key,label in [("name","姓名"),("phone","手机号"),("email","邮箱")] if key != missing_field]
        lines += ["教育：虚构大学计算机本科在读", "技能："+skills, project, experience,
                  "意向：每周可实习四天" if index != 9 else "意向：未说明到岗与实习时间", "声明：所有个人与公司信息为虚构测试数据，禁止真实联系。"]
        (ROOT / f"{sample_id}.txt").write_text("\n".join(lines),encoding="utf-8")
        filename=f"{sample_id}.txt"
        if index==9:
            from docx import Document
            document=Document(); document.add_paragraph('虚构简历：关键字段在表格中')
            for line in lines:
                cells=document.add_table(rows=1,cols=2).rows[0].cells
                label,_,value=line.partition('：'); cells[0].text=label+'：'; cells[1].text=value
            filename=f"{sample_id}.docx"; document.save(ROOT/filename)
        elif index==10:
            filename=f"{sample_id}.pdf"; text_pdf(ROOT/filename,lines)
        samples.append({"id":sample_id,"file":filename,"expected_identity":identity,
                        "manual_correction":{missing_field:identity[missing_field]} if missing_field else {}})
        for job, score_range in zip(JOBS, (backend, frontend)):
            labels.append({"sample_id":sample_id,"job_id":job["id"],"review_status":"pending_human_review", "reviewer":None, "reviewed_at":None,
                "identity":identity,"facts":[project,experience],"requirements":job["requirements"],
                "evidence":[{"file":filename,"quote":project,"limit":"仅支持本人明确说明的项目；技能列表不能证明熟练度"}],
                "missing_information":["缺少聊天核实的到岗意向", "未提及的岗位要求需 HR 核对"],
                "manual_fact":samples[-1]["manual_correction"],"suggested_score_range":score_range,
                "score_range_status":"草案，不能用于宣称质量达标"})
    payload={"version":"demo-v1", "synthetic":True,"samples":samples,"jobs":JOBS,"annotations":labels,
             "ranking_draft":[{"job_id":"backend","higher":"r01","lower":"r08"},{"job_id":"frontend","higher":"r02","lower":"r04"}],
             "quality_status":"pending_human_review"}
    (ROOT / "annotations.json").write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8")
    hashes={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in ROOT.iterdir() if p.is_file()}
    (ROOT / "manifest.json").write_text(json.dumps(hashes,indent=2),encoding="utf-8")
    print(ROOT)


if __name__ == "__main__":
    build()
