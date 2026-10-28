"""剩余验收矩阵：虚构数据、真实本地 API、可控故障与新进程恢复。"""
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
import json
import os
import re
import sqlite3
import subprocess
import sys
from threading import Barrier

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from sqlalchemy import func, select

from tests.test_demo_contracts import env, application, ready_offer, post, enter_interview
import app.main as main
from app.ai.client import LLMClient
from app.ai.service import _mock_evaluation
from app.models import *
from app.services import imports, offers
from app.services.candidates import ingest_candidate
from app.services.evaluations import evaluate_application_record, matching_input
from app.services.periodic import RULES, run_followups
from app.services.workflow import create_application, get_or_create_job_version


def restart(factory, tmp_path):
    """启动另一个 Python 进程；只读此测试数据库和隔离配置，禁止外呼。"""
    database = factory.kw['bind'].url.database
    factory.kw['bind'].dispose()
    environment = {**os.environ, 'APP_DATA_DIR': str(tmp_path), 'APP_RUNTIME_ENV': str(tmp_path/'runtime.env'),
                   'DATABASE_URL': f'sqlite:///{database}', 'LLM_PROVIDER': 'mock', 'SMTP_HOST': '',
                   'WECOM_WEBHOOK_URL': '', 'FOLLOWUP_ENABLED': 'false'}
    code = "import socket; socket.create_connection=lambda *a,**k: (_ for _ in ()).throw(AssertionError('external call')); from app.main import startup; startup()"
    subprocess.run([sys.executable, '-B', '-c', code], env=environment, check=True, capture_output=True, timeout=45)


def upload(client, headers, text, job_id=None, filename='synthetic.txt'):
    response = client.post('/api/import/files', headers=headers, data={'job_id': job_id} if job_id else {},
                           files={'files': (filename, text.encode())})
    assert response.status_code == 200, response.text
    task = response.json()['task']['id']
    parsed = post(client, headers, f'/api/import/tasks/{task}/parse')['task']
    assert parsed['status'] == 'needs_review'
    return task, parsed['preview']['revision']


def paper(client, headers, job_id, kind, suffix=None, title='虚构试卷'):
    files = {}
    if suffix:
        stream = BytesIO()
        if suffix == 'pdf':
            document = PdfWriter(); document.add_blank_page(width=200, height=200); document.write(stream)
        else:
            document = Document(); document.add_paragraph('虚构内部试卷'); document.save(stream)
        files = {'file': ('../../synthetic.'+suffix, stream.getvalue())}
    response = client.post('/api/exams/resources', headers=headers,
        data={'kind': kind, 'title': title, 'job_id': job_id, 'url': 'https://exam.example.test/'+kind, 'request_key': os.urandom(8).hex()}, files=files)
    assert response.status_code == 200, response.text
    return response.json()['resource_id']


def import_sample(env, number=1, job_id=None):
    """冻结样本经正常入口建投递并评估，不直接赋业务状态。"""
    client, headers, factory = env
    samples = Path(__file__).parent/'fixtures'/'demo'
    if job_id is None:
        with factory() as session:
            job, _, _ = get_or_create_job_version(session, {'platform': 'acceptance', 'source_job_id': 'jd-backend',
                'title': '虚构后端岗位', 'jd_raw': (samples/'jd-backend.txt').read_text(encoding='utf-8')})
            session.commit(); job_id = job.id
    task, version = upload(client, headers, (samples/f'r{number:02}.txt').read_text(encoding='utf-8'), job_id)
    result = post(client, headers, f'/api/import/tasks/{task}/confirm', version=version)
    post(client, headers, f"/api/applications/{result['application_id']}/evaluate")
    with factory() as session:
        record = session.get(Application, result['application_id'])
        record.candidate; record.source
        return record


def offer_from_normal_path(env, number=1, *, passed_exams=False):
    client, headers, factory = env
    record = import_sample(env, number)
    if passed_exams:
        post(client, headers, f'/api/applications/{record.id}/decision', status='assessment_pending', reason='HR 审核')
        for version, kind in enumerate(['assessment', 'written']):
            rid = paper(client, headers, record.job_id, kind)
            post(client, headers, f'/api/applications/{record.id}/exams/send', resource_id=rid, candidate_revision=1)
            post(client, headers, f'/api/applications/{record.id}/exams/decision', kind=kind, result='passed', note='HR 已核实', version=version)
    else:
        enter_interview(client, headers, record.id)
    booking = post(client, headers, f'/api/applications/{record.id}/schedule')
    assert client.post(booking['booking_url']+'/confirm', data={'slot_id': booking['slots'][0]['id']}).status_code == 200
    post(client, headers, f"/api/interviews/{booking['interview_id']}/feedback", next_action='offer', result='通过', comment='HR 核实虚构面试证据')
    with factory() as session:
        offer = session.scalar(select(Offer).where(Offer.application_id == record.id))
    values = {'company_name': '虚构验收公司', 'job_title': '虚构后端岗位', 'salary': '虚构300元/天', 'subsidy': '无',
              'office_location': '虚构办公区', 'report_location': '虚构前台', 'report_time': (utc_now()+timedelta(days=7)).date().isoformat(),
              'expires_at': (utc_now()+timedelta(days=3)).isoformat()}
    post(client, headers, f'/api/offers/{offer.id}/terms', version=0, **values)
    post(client, headers, f'/api/offers/{offer.id}/approve', version=1)
    return offer


def test_mixed_batch_extractor_failure_retry_is_local(env, monkeypatch):
    client, headers, factory = env
    monkeypatch.setattr(LLMClient, 'generate_json', lambda *a, **k: pytest.fail('解析不得调用模型'))
    raw = '姓名：虚构批次\n手机号：13900000001\n邮箱：batch@example.test\n项目：Python API'
    result = client.post('/api/import/files', headers=headers, files=[
        ('files', ('good.txt', raw.encode())), ('files', ('broken.pdf', b'broken')),
        ('files', ('transient.txt', raw.replace('批次', '恢复').encode()))]).json()
    ids = [t['id'] for t in result['tasks']]
    original = imports.extract_text_from_file
    def fail_once(path, kind):
        if Path(path).read_bytes() == raw.replace('批次', '恢复').encode():
            raise OSError('synthetic extractor failure')
        return original(path, kind)
    monkeypatch.setattr(imports, 'extract_text_from_file', fail_once)
    states = [post(client, headers, f'/api/import/tasks/{i}/parse')['task']['status'] for i in ids]
    assert states == ['needs_review', 'failed', 'failed']
    confirmed = post(client, headers, f'/api/import/tasks/{ids[0]}/confirm')
    with factory() as session:
        count = session.scalar(select(func.count(ResumeParseRun.id)).where(ResumeParseRun.task_id == ids[0]))
    monkeypatch.setattr(imports, 'extract_text_from_file', original)
    assert post(client, headers, f'/api/import/tasks/{ids[2]}/parse')['task']['status'] == 'needs_review'
    recovered = post(client, headers, f'/api/import/tasks/{ids[2]}/confirm')
    assert recovered['application_id'] != confirmed['application_id']
    assert post(client, headers, f'/api/import/tasks/{ids[0]}/confirm')['application_id'] == confirmed['application_id']
    with factory() as session:
        assert session.scalar(select(func.count(Application.id))) == 2
        assert session.scalar(select(func.count(ResumeParseRun.id)).where(ResumeParseRun.task_id == ids[0])) == count
        assert session.get(ImportTask, ids[1]).status == ImportTaskStatus.FAILED


@pytest.mark.parametrize('leading', ['姓名：虚构规则', '虚构规则'])
@pytest.mark.parametrize('conflicting', [False, True])
def test_rule_parse_unknown_facts_preserved_and_manual_reparse(env, monkeypatch, leading, conflicting):
    client, headers, factory = env
    monkeypatch.setattr(LLMClient, 'generate_json', lambda *a, **k: pytest.fail('规则解析不得调用模型'))
    text = leading+'\n手机号：13900000002\n邮箱：rule@example.test\n经历：跨部门研究；时间交叠待核实，未知工具 FluxLang；没有任职年限证据。'
    if conflicting:
        text += '\n姓名：另一个名字\n手机号：13900000003\n邮箱：conflict@example.test\n备注：以上联系信息归属待 HR 核对'
    task, version = upload(client, headers, text)
    preview = client.get(f'/api/import/tasks/{task}').json()['preview']
    assert preview['candidate']['phone'] == '13900000002'
    # 显式姓名优先于首行猜测；多个字段取首个命中，始终待人工核对。
    assert preview['candidate']['name'] == ('另一个名字' if conflicting and leading == '虚构规则' else '虚构规则')
    assert preview['quality']['warnings']
    assert 'FluxLang' not in preview['candidate'].get('skills', [])
    post(client, headers, f'/api/import/tasks/{task}/supplement', version=version, name='虚构规则', skills='FluxLang', experience='HR 核实：内部研究协作，未确认年限')
    parsed = post(client, headers, f'/api/import/tasks/{task}/parse')['task']['preview']
    assert parsed['candidate']['skills'] == ['FluxLang']
    result = post(client, headers, f'/api/import/tasks/{task}/confirm', version=parsed['revision'])
    with factory() as session:
        record = session.get(Application, result['application_id'])
        assert record.source.raw_resume == text
        assert 'FluxLang' in json.dumps(record.candidate.structured_data)
        assert session.scalar(select(func.count(Evaluation.id))) == 0


@pytest.mark.parametrize('status', [ApplicationStatus.NEXT_ROUND, ApplicationStatus.OFFER_PENDING])
def test_api_concurrent_import_preserves_late_stage(env, status):
    client, headers, factory = env
    with factory() as session:
        record = application(session, status)
        raw = record.source.raw_resume
        job_id, record_id = record.job_id, record.id
        if status == ApplicationStatus.NEXT_ROUND:
            session.add(Interview(application_id=record.id, round=1, status=InterviewStatus.COMPLETED, feedback={'next_action': 'next_round'}))
            session.commit()
    tasks = [upload(client, headers, raw, job_id) for _ in range(2)]
    barrier = Barrier(2)
    def confirm(pair):
        barrier.wait(timeout=10)
        return post(client, headers, f'/api/import/tasks/{pair[0]}/confirm', version=pair[1])
    with ThreadPoolExecutor(2) as pool:
        values = list(pool.map(confirm, tasks))
    assert [v['application_id'] for v in values] == [record_id]*2
    for task, version in tasks:
        assert post(client, headers, f'/api/import/tasks/{task}/confirm', version=version)['application_id'] == record_id
    with factory() as session:
        record = session.get(Application, record_id)
        assert record.status == status and len(record.source_links) == 3
        assert session.scalar(select(func.count(Candidate.id))) == 1
        assert session.scalar(select(func.count(Application.id))) == 1
        job, jd, _ = get_or_create_job_version(session, {'platform': 'test', 'source_job_id': 'same-title-new-id', 'title': record.job.title, 'jd_raw': 'Rust'})
        session.commit(); other_job_id = job.id
    task, version = upload(client, headers, raw, other_job_id)
    assert post(client, headers, f'/api/import/tasks/{task}/confirm', version=version)['application_id'] != record_id


@pytest.mark.parametrize('failure', ['401', '429', 'timeout', 'json', 'schema'])
def test_online_fault_matrix_retains_old_evaluation_and_retries(env, monkeypatch, failure):
    client, headers, factory = env
    with factory() as session:
        record = application(session, ApplicationStatus.OFFER_PENDING)
        old = evaluate_application_record(session, record)
        facts = deepcopy(record.candidate.structured_data)
        snapshot = matching_input(record)
    payload = _mock_evaluation(snapshot)
    for dimension in payload['dimensions']:
        dimension.update(score=0, references=[])
    for key, value in {'LLM_PROVIDER': 'openai_compatible', 'LLM_BASE_URL': 'https://model.example.test/v1', 'LLM_MODEL': 'synthetic', 'LLM_API_KEY': 'synthetic-only'}.items():
        monkeypatch.setenv(key, value)
    calls = []
    def request(url, **kwargs):
        calls.append(url)
        if failure == 'timeout': raise httpx.ReadTimeout('synthetic timeout')
        if failure in {'401', '429'}: return httpx.Response(int(failure), request=httpx.Request('POST', url))
        return httpx.Response(200, json={'choices': [{'message': {'content': 'bad json' if failure == 'json' else '{}'}}]}, request=httpx.Request('POST', url))
    monkeypatch.setattr('app.ai.client.httpx.post', request)
    monkeypatch.setattr('app.ai.service._mock_evaluation', lambda *a, **k: pytest.fail('在线失败不得回退 Mock'))
    failed = client.post(f'/api/applications/{record.id}/evaluate', headers=headers)
    assert failed.status_code >= 400 and len(calls) == 1
    with factory() as session:
        assert session.scalar(select(func.count(Evaluation.id))) == 1
        assert session.get(Evaluation, old['evaluation_id']).is_stale is False
    monkeypatch.setattr('app.ai.client.httpx.post', lambda url, **kwargs: httpx.Response(200, json={'choices': [{'message': {'content': json.dumps(payload)}}]}, request=httpx.Request('POST', url)))
    post(client, headers, f'/api/applications/{record.id}/evaluate')
    with factory() as session:
        assert session.scalar(select(func.count(Evaluation.id))) == 2
        assert session.scalars(select(EvaluationAttempt.status).order_by(EvaluationAttempt.id)).all() == ['succeeded', 'failed', 'succeeded']
        assert session.get(Application, record.id).status == ApplicationStatus.OFFER_PENDING
        assert session.get(Candidate, record.candidate_id).structured_data == facts


@pytest.mark.parametrize('rule', RULES, ids=lambda r: r.code)
def test_followup_each_threshold_persists_without_page(env, rule):
    _, _, factory = env
    now = datetime(2026, 11, 2, 2, tzinfo=timezone.utc)
    status = 'interviewed' if rule.code == 'interview_feedback_pending' else rule.statuses[0]
    with factory() as session:
        record = application(session, ApplicationStatus(status))
        moment = now - rule.threshold
        record.created_at = record.applied_at = record.status_changed_at = moment
        if rule.mode == 'upcoming':
            # 日期字段按北京时间零点解释。
            now = datetime(2026, 11, 1, 16, tzinfo=timezone.utc)
            record.extra_data = {'onboarding': {'expected_date': '2026-11-05'}}
        session.commit()
    def count():
        with factory() as session:
            return session.scalar(select(func.count(FollowUpTask.id)).where(FollowUpTask.rule_code == rule.code, FollowUpTask.target_id == record.id))
    run_followups(factory, now=now-timedelta(seconds=1)); assert count() == 0
    run_followups(factory, now=now+timedelta(seconds=1)); assert count() == 1
    run_followups(factory, now=now+timedelta(seconds=2)); assert count() == 1


@pytest.mark.parametrize('kind', ['assessment', 'written'])
@pytest.mark.parametrize('suffix', [None, 'pdf', 'docx'])
def test_resource_formats_isolation_and_unused_delete(env, kind, suffix):
    client, headers, factory = env
    with factory() as session:
        record = application(session)
    rid = paper(client, headers, record.job_id, kind, suffix)
    with factory() as session:
        resource = session.get(ExamResource, rid)
        if suffix:
            assert Path(resource.file_path).resolve().is_relative_to(imports.UPLOAD_DIR.resolve())
            content = Path(resource.file_path).read_bytes()
            assert client.get(f'/api/exams/resources/{rid}/file').content == content
    with TestClient(main.app) as anonymous:
        assert anonymous.get(f'/api/exams/resources/{rid}/file').status_code == 401
        assert anonymous.post(f'/api/exams/resources/{rid}/delete').status_code in {401, 403}
    post(client, headers, f'/api/exams/resources/{rid}/delete')
    with factory() as session:
        assert session.get(ExamResource, rid) is None
    if suffix: assert not Path(resource.file_path).exists()


@pytest.mark.parametrize('kind', ['assessment', 'written'])
def test_resource_validation_matrix_no_partial_writes(env, monkeypatch, kind):
    client, headers, factory = env
    with factory() as session: record = application(session)
    good = {'job_id': record.job_id, 'kind': kind, 'title': '虚构试卷', 'url': 'https://exam.example.test/paper'}
    bad = [{k: v for k, v in good.items() if k != missing} for missing in ['job_id', 'title', 'url']]
    bad += [{**good, 'url': url} for url in ['javascript:alert(1)', 'file:///tmp/x', 'https://user:password@example.test', 'https://exa mple.test']]
    bad += [{**good, 'job_id': -1}, {**good, 'kind': 'unknown'}]
    for values in bad:
        assert client.post('/api/exams/resources', headers=headers, data=values).status_code == 422
    monkeypatch.setenv('MAX_UPLOAD_BYTES', '20')
    for name, content in [('empty.pdf', b''), ('wrong.pdf', b'bad'), ('wrong.docx', b'bad'), ('wrong.exe', b'x'), ('large.pdf', b'%PDF-'+b'x'*20)]:
        assert client.post('/api/exams/resources', headers=headers, data=good, files={'file': (name, content)}).status_code == 422
    with factory() as session: assert session.scalar(select(func.count(ExamResource.id))) == 0


@pytest.mark.parametrize('kind', ['assessment', 'written'])
@pytest.mark.parametrize('outcome', ['sent', 'failed', 'unknown', 'sending'])
def test_exam_failure_restart_retry_disable_and_snapshot(env, monkeypatch, tmp_path, kind, outcome):
    client, headers, factory = env
    record = import_sample(env)
    recipient = record.candidate.email
    post(client, headers, f'/api/applications/{record.id}/decision', status='assessment_pending', reason='HR 核实')
    if kind == 'written':
        post(client, headers, f'/api/applications/{record.id}/exams/decision', kind='assessment', result='skipped', note='HR 免试', version=0)
    rid = paper(client, headers, record.job_id, kind, 'pdf')
    calls = []
    class SMTP:
        def __init__(self, *a, **k):
            if outcome == 'failed': raise ConnectionRefusedError('synthetic')
        def __enter__(self): return self
        def __exit__(self, *a): pass
        def starttls(self): pass
        def login(self, *a): pass
        def send_message(self, message):
            calls.append(message)
            assert not list(message.iter_attachments())
            assert message['To'] == recipient
            if outcome == 'unknown': raise TimeoutError('synthetic after DATA')
    monkeypatch.setenv('SMTP_HOST', 'smtp.example.test'); monkeypatch.setenv('SMTP_SSL', 'false')
    monkeypatch.setattr(offers.smtplib, 'SMTP', SMTP)
    endpoint = f'/api/applications/{record.id}/exams/send'
    result = post(client, headers, endpoint, resource_id=rid, candidate_revision=1)
    nid = result['notification_id']
    assert result['status'] == ('sent' if outcome == 'sending' else outcome)
    with factory() as session:
        notification = session.get(Notification, nid)
        dispatch = session.get(ExamDispatch, notification.target_id)
        snapshot = deepcopy(dispatch.snapshot)
        filepath = Path(session.get(ExamResource, rid).file_path); content = filepath.read_bytes()
        if outcome == 'sending': notification.status = NotificationStatus.SENDING
        session.commit()
    restart(factory, tmp_path)
    with factory() as session:
        expected = 'unknown' if outcome == 'sending' else outcome
        assert session.get(Notification, nid).status.value == expected
        assert session.get(ExamDispatch, dispatch.id).snapshot == snapshot
        assert filepath.read_bytes() == content
    if outcome in {'unknown', 'sending', 'sent'}:
        repeated = post(client, headers, endpoint, resource_id=rid, candidate_revision=1)
        assert repeated['notification_id'] == nid and len(calls) == (0 if outcome == 'failed' else 1)
    if outcome in {'unknown', 'sending'}:
        post(client, headers, f'/api/notifications/{nid}/verify', delivered='false', note='HR 核对专用服务日志，确认未送达')
    post(client, headers, f'/api/exams/resources/{rid}/delete')
    assert client.post(endpoint, headers=headers, data={'resource_id': rid, 'candidate_revision': 1}).status_code == 409
    with factory() as session:
        assert not session.get(ExamResource, rid).active
        assert session.get(ExamDispatch, dispatch.id).snapshot == snapshot
    if outcome != 'sent':
        new = paper(client, headers, record.job_id, kind, 'docx', title='更换虚构试卷')
        monkeypatch.setenv('SMTP_HOST', '')
        assert post(client, headers, endpoint, resource_id=new, candidate_revision=1)['status'] == 'mocked'
        with factory() as session:
            assert session.scalar(select(func.count(ExamDispatch.id))) == 2
            assert session.get(ExamDispatch, dispatch.id).snapshot == snapshot
        post(client, headers, f'/api/applications/{record.id}/exams/decision', kind=kind, result='passed', note='HR 核实', version=0 if kind == 'assessment' else 1)


@pytest.mark.parametrize('failure', ['timeout', 'refused'])
def test_meeting_failure_is_traced_and_same_invite_can_retry(env, monkeypatch, failure):
    client, headers, factory = env
    with factory() as session: record = application(session)
    enter_interview(client, headers, record.id)
    schedule = post(client, headers, f'/api/applications/{record.id}/schedule')
    original = main.create_mock_tencent_meeting
    def failing(**kwargs): raise (TimeoutError if failure == 'timeout' else ValueError)('synthetic meeting failure')
    monkeypatch.setattr(main, 'create_mock_tencent_meeting', failing)
    response = client.post(schedule['booking_url']+'/confirm', data={'slot_id': schedule['slots'][0]['id']})
    assert response.status_code == 503
    with factory() as session:
        interview = session.get(Interview, schedule['interview_id'])
        assert interview.confirmed_start_at is None and interview.booking_token
        assert interview.meeting_status == 'mock_failed'
        assert session.scalar(select(AuditLog).where(AuditLog.action == 'interview.meeting_failed'))
    monkeypatch.setattr(main, 'create_mock_tencent_meeting', original)
    assert client.post(schedule['booking_url']+'/confirm', data={'slot_id': schedule['slots'][0]['id']}).status_code == 200


@pytest.mark.parametrize('failure', ['timeout', 'refused'])
def test_wecom_failure_persists_and_explicit_retry_recovers(env, monkeypatch, failure):
    client, headers, factory = env
    monkeypatch.setenv('WECOM_WEBHOOK_URL', 'https://wecom.example.test/synthetic')
    def request(url, **kwargs):
        if failure == 'timeout': raise httpx.ReadTimeout('synthetic timeout')
        return httpx.Response(200, json={'errcode': 40001, 'errmsg': 'synthetic rejected'}, request=httpx.Request('POST', url))
    monkeypatch.setattr('app.adapters.notifications.httpx.post', request)
    with factory() as session:
        record = application(session)
        notification = offers.notify_internal(session, title='虚构内部提醒', content='需要 HR 核实', target_type='application', target_id=record.id)
        session.commit(); nid = notification.id
        assert notification.status == NotificationStatus.FAILED and notification.error_message
    monkeypatch.setattr('app.adapters.notifications.httpx.post', lambda url, **kwargs: httpx.Response(200, json={'errcode': 0}, request=httpx.Request('POST', url)))
    result = post(client, headers, f'/api/notifications/{nid}/retry')
    assert result['status'] == 'sent' and result['message'] == '内部通知已发送'
    assert post(client, headers, f'/api/notifications/{nid}/retry')['notification_id'] == result['notification_id']
    with factory() as session:
        assert session.get(Notification, nid).status == NotificationStatus.FAILED
        assert session.scalar(select(func.count(Notification.id))) == 2
        assert session.get(Application, record.id).status == ApplicationStatus.EVALUATION_REVIEW


def login(client, username):
    token = re.search(r'name="csrf_token" value="([^"]+)"', client.get('/register').text)[1]
    assert client.post('/register', data={'username': username, 'password': 'password123', 'csrf_token': token}).status_code == 200
    token = re.search(r'name="csrf_token" value="([^"]+)"', client.get('/login').text)[1]
    response = client.post('/login', data={'username': username, 'password': 'password123', 'csrf_token': token})
    return {'X-CSRF-Token': re.search(r'name="csrf-token" content="([^"]+)"', response.text)[1]}


def test_both_hr_users_all_business_audit_actions(env):
    first, headers, factory = env
    with TestClient(main.app) as second:
        second_headers = login(second, 'second-hr')
        for client, csrf, name in [(first, headers, 'demo-hr'), (second, second_headers, 'second-hr')]:
            with factory() as session:
                record = application(session)
                offer, values = ready_offer(session)
                start = session.scalar(select(func.max(AuditLog.id))) or 0
            post(client, csrf, f'/api/candidates/{record.candidate_id}/profile', version=1, skills='Python, Rust', note='HR 已核实')
            post(client, csrf, f'/api/applications/{record.id}/decision', status='assessment_pending', reason='HR 核实')
            for version, kind in enumerate(['assessment', 'written']):
                rid = paper(client, csrf, record.job_id, kind)
                post(client, csrf, f'/api/applications/{record.id}/exams/send', resource_id=rid, candidate_revision=2)
                post(client, csrf, f'/api/applications/{record.id}/exams/decision', kind=kind, result='passed', note='HR 已核实', version=version)
                post(client, csrf, f'/api/exams/resources/{rid}/delete')
            post(client, csrf, f'/api/offers/{offer.id}/return', version=1, reason='修订薪资')
            post(client, csrf, f'/api/offers/{offer.id}/terms', version=1, **{**values, 'salary': '虚构：350 元/天'})
            post(client, csrf, f'/api/offers/{offer.id}/approve', version=2)
            post(client, csrf, f'/api/offers/{offer.id}/send')
            with factory() as session: token = session.get(Offer, offer.id).candidate_token
            with TestClient(main.app) as candidate:
                assert candidate.post(f'/offer/{token}/respond', data={'accepted': 'true'}).status_code == 200
            post(client, csrf, f'/api/applications/{offer.application_id}/onboarding', action='update', version=0,
                 expected_date=utc_now().date().isoformat(), location='虚构前台', note='HR 核实登记')
            with factory() as session:
                logs = session.scalars(select(AuditLog).where(AuditLog.id > start)).all()
                required = ['candidate.supplemented', 'application.status_changed', 'exam_resource.created', 'exam.dispatched',
                            'exam.assessment_decided', 'exam.written_decided', 'offer.approved', 'onboarding.update']
                for action in required:
                    assert any(log.action == action and log.actor.endswith(':'+name) for log in logs), action
                assert any(log.actor == '候选人' for log in logs)


def test_every_internal_route_auth_and_csrf_no_writes(env):
    client, headers, factory = env
    with factory() as session: record = application(session)
    database = factory.kw['bind'].url.database
    def snapshot():
        with sqlite3.connect(database) as connection: return '\n'.join(connection.iterdump())
    before = snapshot()
    routes = [(re.sub(r'\{[^}]+\}', '1', r.path), r.methods) for r in main.app.routes if getattr(r, 'path', '').startswith('/api/')]
    with TestClient(main.app) as anonymous:
        for path, methods in routes:
            for method in methods & {'GET', 'POST'}:
                assert anonymous.request(method, path, headers={'X-CSRF-Token': 'forged'}).status_code == (401 if method == 'GET' else 403), path
        for path in ['/login', '/register']:
            for token in ['', 'forged']:
                assert anonymous.post(path, data={'username': 'forged-user', 'password': 'password123', 'csrf_token': token}).status_code == 403
    for path, methods in routes:
        if 'POST' in methods:
            assert client.post(path).status_code == 403, path
            assert client.post(path, headers={'X-CSRF-Token': 'forged'}).status_code == 403, path
    for path in ['/settings', '/settings/wecom', '/settings/llm/clear']:
        assert client.post(path, data={'provider': 'mock', 'csrf_token': 'forged'}).status_code == 403
    assert snapshot() == before
    post(client, headers, f'/api/applications/{record.id}/decision', status='on_hold', reason='合法 CSRF')


def test_public_token_cannot_choose_another_invite_or_read_internal(env):
    client, headers, factory = env
    with factory() as session: records = [application(session) for _ in range(2)]
    invitations = []
    for record in records:
        enter_interview(client, headers, record.id)
        invitations.append(post(client, headers, f'/api/applications/{record.id}/schedule'))
    with TestClient(main.app) as anonymous:
        response = anonymous.post(invitations[0]['booking_url']+'/confirm', data={'slot_id': invitations[1]['slots'][0]['id']})
        assert response.status_code == 404
        assert anonymous.get('/api/candidates').status_code == 401
        assert anonymous.post('/schedule/forged/confirm', data={'slot_id': invitations[0]['slots'][0]['id']}).status_code == 400
        assert anonymous.post('/offer/forged/respond', data={'accepted': 'true'}).status_code == 422
    with factory() as session: assert all(session.get(Interview, i['interview_id']).confirmed_start_at is None for i in invitations)


@pytest.mark.parametrize('change', [{'name': '同名', 'phone': '13900000012'}, {'name': '不同姓名', 'phone': '13900000011'}])
def test_weak_identity_never_merges_and_hr_can_correct(env, change):
    client, headers, factory = env
    with factory() as session:
        existing, _, _ = ingest_candidate(session, candidate_data={'name': '同名', 'phone': '13900000011', 'email': 'original@example.test'},
                                         source_data={'platform': 'test', 'source_application_id': 'original'})
        legacy = Candidate(name='历史待补', normalized_name='历史待补', dedup_status=DedupStatus.NEEDS_REVIEW)
        session.add(legacy); session.commit()
    raw = f"姓名：{change['name']}\n手机号：{change['phone']}\n邮箱：new@example.test\n项目：Python"
    task, version = upload(client, headers, raw)
    result = post(client, headers, f'/api/import/tasks/{task}/confirm', version=version)
    with factory() as session:
        candidate = session.get(Application, result['application_id']).candidate
        cid = candidate.id
        assert cid != existing.id and candidate.dedup_status.value == 'needs_review'
        assert '未自动合并' in str(candidate.structured_data['quality']['warnings'])
        assert session.scalar(select(func.count(Candidate.id))) == 3
        assert session.get(Candidate, legacy.id).email is None
    post(client, headers, f'/api/candidates/{cid}/profile', version=1, name='虚构核实新身份', phone='13900000013', note='HR 确认为不同候选人')
    assert client.post(f'/api/candidates/{legacy.id}/profile', headers=headers, data={'note': '未补身份'}).status_code in {409, 422}
    with factory() as session:
        assert session.scalar(select(AuditLog).where(AuditLog.action == 'candidate.supplemented', AuditLog.entity_id == cid)).actor.endswith(':demo-hr')


def test_schedule_fixed_clock_hours_buffer_and_confirmation_recheck(env, monkeypatch):
    from zoneinfo import ZoneInfo
    client, headers, factory = env
    now = datetime(2026, 11, 6, 13, tzinfo=timezone.utc)  # 周五北京时间 21:00
    class Clock(datetime):
        @classmethod
        def now(cls, tz=None): return now.astimezone(tz) if tz else now.replace(tzinfo=None)
    monkeypatch.setattr(main, 'datetime', Clock)
    with factory() as session:
        record = application(session)
        busy = application(session, ApplicationStatus.INTERVIEW_CONFIRMED)
        occupied = Interview(application_id=busy.id, round=1, status=InterviewStatus.CONFIRMED,
                             confirmed_start_at=datetime(2026, 11, 9, 9, tzinfo=ZoneInfo('Asia/Shanghai')), confirmed_end_at=datetime(2026, 11, 9, 10, tzinfo=ZoneInfo('Asia/Shanghai')))
        session.add(occupied); session.commit()
    assert client.post(f'/api/applications/{record.id}/schedule', headers=headers).status_code == 409
    enter_interview(client, headers, record.id)
    schedule = post(client, headers, f'/api/applications/{record.id}/schedule')
    for slot in schedule['slots']:
        # 数据库的预约字段存储本地带时区的墙钟值，appointment_time 负责解释。
        start = main.appointment_time(datetime.fromisoformat(slot['start_at'])).astimezone(ZoneInfo('Asia/Shanghai'))
        end = main.appointment_time(datetime.fromisoformat(slot['end_at'])).astimezone(ZoneInfo('Asia/Shanghai'))
        assert start > now and start.weekday() < 5 and start.hour >= 9 and end.hour <= 18
        assert end-start == timedelta(hours=1)
        assert not main.intervals_conflict(start, end, occupied.confirmed_start_at, occupied.confirmed_end_at, buffer_minutes=15)
    with factory() as session:
        slot = session.get(InterviewSlot, schedule['slots'][0]['id'])
        slot.start_at, slot.end_at = now-timedelta(hours=2), now-timedelta(hours=1)
        session.commit()
    assert client.post(schedule['booking_url']+'/confirm', data={'slot_id': slot.id}).status_code == 409


@pytest.mark.parametrize('stage', ['screen', 'interview'])
def test_rejection_never_changes_same_person_other_job(env, stage):
    client, headers, factory = env
    with factory() as session:
        record = application(session)
        job, jd, _ = get_or_create_job_version(session, {'platform': 'test', 'source_job_id': 'other', 'title': '其他岗位', 'jd_raw': 'Rust'})
        other, _ = create_application(session, {'candidate_id': record.candidate_id, 'job_id': job.id, 'job_version_id': jd.id})
        session.commit(); other_status = other.status
    if stage == 'screen': post(client, headers, f'/api/applications/{record.id}/decision', status='rejected', reason='本岗位不匹配')
    else:
        enter_interview(client, headers, record.id)
        schedule = post(client, headers, f'/api/applications/{record.id}/schedule')
        assert client.post(schedule['booking_url']+'/confirm', data={'slot_id': schedule['slots'][0]['id']}).status_code == 200
        post(client, headers, f"/api/interviews/{schedule['interview_id']}/feedback", next_action='reject', result='未通过', comment='本岗位不匹配')
        assert client.post(f"/api/interviews/{schedule['interview_id']}/send-invite", headers=headers).status_code == 409
    for action in ['schedule', 'offer']:
        assert client.post(f'/api/applications/{record.id}/{action}', headers=headers).status_code == (409 if action == 'schedule' else 422)
    with factory() as session:
        assert session.get(Application, record.id).status == ApplicationStatus.REJECTED
        assert session.get(Application, other.id).status == other_status


@pytest.mark.parametrize('branch', ['needs_supplement', 'defer', 'withdraw'])
def test_onboarding_branches_restart_reminders_and_frozen_offer(env, tmp_path, branch):
    client, headers, factory = env
    prepared = offer_from_normal_path(env)
    with factory() as session:
        offer = session.get(Offer, prepared.id)
        offers.send_offer(session, offer.id); session.commit()
        offers.respond_offer(session, offer.candidate_token, accepted=True); session.commit()
        snapshot = offers.offer_snapshot(offer)
        tasks = session.scalars(select(OnboardingTask).where(OnboardingTask.application_id == offer.application_id)).all()
    endpoint = f'/api/applications/{offer.application_id}/onboarding'
    today = utc_now().date().isoformat()
    post(client, headers, endpoint, action='update', version=0, expected_date=today, location='虚构前台')
    run_followups(factory)
    post(client, headers, f'/api/onboarding/{tasks[0].id}/complete')
    post(client, headers, endpoint, action=branch, version=1, expected_date=(utc_now()+timedelta(days=10)).date().isoformat(), note='HR 已核实分支原因')
    run_followups(factory)
    restart(factory, tmp_path)
    with factory() as session:
        assert offers.offer_snapshot(session.get(Offer, offer.id)) == snapshot
        assert session.get(Offer, offer.id).status == OfferStatus.ACCEPTED
        assert session.get(OnboardingTask, tasks[0].id).status == OnboardingStatus.DONE
        if branch in {'defer', 'withdraw'}:
            assert not session.scalar(select(FollowUpTask.id).where(FollowUpTask.target_type == 'application', FollowUpTask.target_id == offer.application_id,
                    FollowUpTask.rule_code == 'onboarding_approaching', FollowUpTask.status == FollowUpStatus.OPEN))
    if branch == 'withdraw':
        assert client.post(f'/api/onboarding/{tasks[1].id}/complete', headers=headers).status_code == 422
        assert client.post(endpoint, headers=headers, data={'action': 'join', 'version': 2, 'actual_date': today, 'note': '旧请求'}).status_code == 409
        assert client.get('/api/dashboard/summary').json()['status_counts']['onboarding_withdrawn'] == 1
        return
    if branch == 'needs_supplement':
        assert client.post(endpoint, headers=headers, data={'action': 'join', 'version': 2, 'actual_date': today, 'note': '材料未齐'}).status_code == 409
    post(client, headers, endpoint, action='update', version=2, expected_date=today, location='虚构前台', note='补齐后继续')
    for task in tasks: post(client, headers, f'/api/onboarding/{task.id}/complete')
    post(client, headers, endpoint, action='join', version=3, actual_date=today, note='HR 现场核实')
    restart(factory, tmp_path)
    post(client, headers, endpoint, action='join', version=3, actual_date=today, note='HR 现场核实')
    with factory() as session:
        assert session.scalar(select(func.count(AuditLog.id)).where(AuditLog.action == 'onboarding.join', AuditLog.entity_id == offer.application_id)) == 1
        assert session.get(Application, offer.application_id).status == ApplicationStatus.JOINED


def test_beijing_midnight_dashboard_exports_and_stage_filters(env, monkeypatch):
    import csv
    from io import StringIO
    from openpyxl import load_workbook
    client, headers, factory = env
    now = datetime(2026, 11, 2, 15, 59, 59, tzinfo=timezone.utc)
    class Clock(datetime):
        @classmethod
        def now(cls, tz=None): return now.astimezone(tz) if tz else now.replace(tzinfo=None)
    monkeypatch.setattr(main, 'datetime', Clock)
    expected = {}
    for decision in ['passed', 'skipped', 'failed']:
        with factory() as session:
            record = application(session)
            name = record.candidate.name
        post(client, headers, f'/api/applications/{record.id}/decision', status='assessment_pending', reason='HR 审核')
        rid = paper(client, headers, record.job_id, 'assessment')
        post(client, headers, f'/api/applications/{record.id}/exams/send', resource_id=rid, candidate_revision=1)
        post(client, headers, f'/api/applications/{record.id}/exams/decision', kind='assessment', result=decision, note='HR 核实', version=0)
        if decision != 'failed':
            post(client, headers, f'/api/applications/{record.id}/exams/decision', kind='written', result='skipped', note='HR 免试', version=1)
        expected[name] = (decision, 'skipped' if decision != 'failed' else '历史未记录 / 未办理', '', '')
        with factory() as session:
            row = session.get(Application, record.id)
            row.created_at = now-timedelta(days=10); row.updated_at = now; session.commit()
    with factory() as session:
        offer, _ = ready_offer(session)
        offers.send_offer(session, offer.id); session.commit()
        offers.respond_offer(session, offer.candidate_token, accepted=False); session.commit()
        offer.application.created_at = now-timedelta(days=10)
        expected[offer.candidate_name] = ('历史未记录 / 未办理', '历史未记录 / 未办理', '', '')
        for index, state in enumerate([ApplicationStatus.ONBOARDING, ApplicationStatus.JOINED]):
            row = application(session, state)
            row.created_at = now if index == 0 else now+timedelta(seconds=1)
            row.extra_data = {'onboarding': {'state': 'preparing' if index == 0 else 'joined', 'actual_date': '' if index == 0 else '2026-11-03'}}
            expected[row.candidate.name] = ('历史未记录 / 未办理', '历史未记录 / 未办理', 'preparing' if index == 0 else 'joined', '' if index == 0 else '2026-11-03')
        session.commit()
    before = client.get('/api/dashboard/summary').json()
    assert before['new_today'] == 1
    now += timedelta(seconds=1)
    after = client.get('/api/dashboard/summary').json()
    assert after['new_today'] == 1 and after['funnel']['offers'] == 1
    assert after['exam_counts']['assessment'] == {'sent': 3, 'passed': 1, 'failed': 1, 'skipped': 1, 'historical_unrecorded': 3}
    assert after['exam_counts']['written'] == {'sent': 0, 'passed': 0, 'failed': 0, 'skipped': 2, 'historical_unrecorded': 3}
    csv_rows = list(csv.reader(StringIO(client.get('/api/exports/recruitment.csv').content.decode('utf-8-sig'))))
    workbook = load_workbook(BytesIO(client.get('/api/exports/recruitment.xlsx').content), read_only=True)
    xlsx_rows = [[str(value) if value is not None else '' for value in row] for row in workbook.active.values]
    assert csv_rows == xlsx_rows
    assert csv_rows[0] == ['候选人', '岗位', '状态', '测评结论', '笔试结论', '入职办理', '实际到岗日期', '评分', '负责人', '面试时间', '更新时间']
    assert {row[0]: tuple(row[3:7]) for row in csv_rows[1:]} == expected
    rows = client.get('/api/candidates').json()
    assert sum(row['status'] == 'joined' for row in rows) == after['status_counts']['joined'] == 1
    html = client.get('/candidates').text
    assert '测评' in html and '笔试' in html and '已入职' in html


@pytest.mark.parametrize('kind', ['assessment', 'written'])
def test_exam_same_resource_definite_failure_retry(env, monkeypatch, kind):
    client, headers, factory = env
    with factory() as session: record = application(session, ApplicationStatus.ASSESSMENT_PENDING if kind == 'assessment' else ApplicationStatus.WRITTEN_PENDING)
    rid = paper(client, headers, record.job_id, kind)
    endpoint = f'/api/applications/{record.id}/exams/send'
    monkeypatch.setenv('SMTP_HOST', 'smtp.example.test')
    monkeypatch.setattr(offers.smtplib, 'SMTP', lambda *a, **k: (_ for _ in ()).throw(ConnectionRefusedError('synthetic')))
    first = post(client, headers, endpoint, resource_id=rid, candidate_revision=1)
    assert first['status'] == 'failed'
    monkeypatch.setenv('SMTP_HOST', '')
    second = post(client, headers, endpoint, resource_id=rid, candidate_revision=1)
    assert second['status'] == 'mocked' and second['notification_id'] != first['notification_id']
    assert post(client, headers, endpoint, resource_id=rid, candidate_revision=1)['notification_id'] == second['notification_id']
    with factory() as session:
        assert session.scalar(select(func.count(ExamDispatch.id))) == 1
        assert session.scalar(select(func.count(Notification.id))) == 2


@pytest.mark.parametrize('kind', ['assessment', 'written'])
@pytest.mark.parametrize('decision', ['passed', 'failed', 'skipped'])
def test_s07_s08_from_import_exam_branches_and_cross_resource_gate(env, kind, decision):
    client, headers, factory = env
    record = import_sample(env, 7 if kind == 'assessment' else 8)
    post(client, headers, f'/api/applications/{record.id}/decision', status='assessment_pending', reason='HR 审核')
    version = 0
    if kind == 'written':
        post(client, headers, f'/api/applications/{record.id}/exams/decision', kind='assessment', result='skipped', note='HR 免测评', version=0)
        version = 1
    endpoint = f'/api/applications/{record.id}/exams/send'
    opposite = paper(client, headers, record.job_id, 'written' if kind == 'assessment' else 'assessment')
    assert client.post(endpoint, headers=headers, data={'resource_id': opposite, 'candidate_revision': 1}).status_code == 409
    with factory() as session:
        job, _, _ = get_or_create_job_version(session, {'platform': 'acceptance', 'source_job_id': 'jd-frontend', 'title': '虚构前端岗位', 'jd_raw': 'Vue TypeScript'})
        session.commit()
    wrong_job = paper(client, headers, job.id, kind)
    assert client.post(endpoint, headers=headers, data={'resource_id': wrong_job, 'candidate_revision': 1}).status_code == 409
    if decision != 'skipped':
        rid = paper(client, headers, record.job_id, kind)
        post(client, headers, endpoint, resource_id=rid, candidate_revision=1)
    assert client.post(f'/api/applications/{record.id}/schedule', headers=headers).status_code == 409
    values = {'kind': kind, 'result': decision, 'note': 'HR 外部核实或有理由免试', 'version': version}
    result = post(client, headers, f'/api/applications/{record.id}/exams/decision', **values)
    expected = 'rejected' if decision == 'failed' else 'written_pending' if kind == 'assessment' else 'interview_pending'
    assert result['status'] == expected
    assert post(client, headers, f'/api/applications/{record.id}/exams/decision', **values)['status'] == expected
    assert client.post(f'/api/applications/{record.id}/exams/decision', headers=headers, data={**values, 'result': 'skipped' if decision != 'skipped' else 'failed'}).status_code == 409


@pytest.mark.parametrize('outcome', ['declined', 'expired'])
def test_s03_offer_branches_from_import(env, monkeypatch, outcome):
    client, headers, factory = env
    offer = offer_from_normal_path(env, 3)
    post(client, headers, f'/api/offers/{offer.id}/send')
    with factory() as session: token = session.get(Offer, offer.id).candidate_token
    if outcome == 'expired': monkeypatch.setattr(offers, 'utc_now', lambda: utc_now()+timedelta(days=10))
    response = client.post(f'/offer/{token}/respond', data={'accepted': 'false' if outcome == 'declined' else 'true'})
    assert response.status_code == 200
    with factory() as session: assert session.get(Offer, offer.id).status.value == outcome
    assert client.post(f'/offer/{token}/respond', data={'accepted': 'true'}).status_code == 422
    with factory() as session: assert session.scalar(select(func.count(OnboardingTask.id))) == 0


@pytest.mark.parametrize('target', ['next_round', 'offer_pending'])
def test_s05_from_import_repeat_source_manual_facts_and_reassessment(env, target):
    client, headers, factory = env
    record = import_sample(env, 5)
    enter_interview(client, headers, record.id)
    booking = post(client, headers, f'/api/applications/{record.id}/schedule')
    assert client.post(booking['booking_url']+'/confirm', data={'slot_id': booking['slots'][0]['id']}).status_code == 200
    post(client, headers, f"/api/interviews/{booking['interview_id']}/feedback", next_action='next_round' if target == 'next_round' else 'offer', result='通过', comment='虚构核实')
    duplicate = import_sample(env, 5, record.job_id)
    assert duplicate.id == record.id and duplicate.status.value == target
    post(client, headers, f'/api/candidates/{record.candidate_id}/profile', version=1, experience='HR 核实的新 Rust 项目', skills='Rust', application_id=record.id, chat_summary='仅此岗位新增沟通')
    with factory() as session:
        old = session.scalars(select(Evaluation).where(Evaluation.application_id == record.id)).all()
        assert all(e.is_stale for e in old)
    post(client, headers, f'/api/applications/{record.id}/evaluate')
    with factory() as session:
        current = session.get(Application, record.id)
        assert current.status.value == target and len(current.source_links) == 2
        latest = session.scalar(select(Evaluation).where(Evaluation.application_id == record.id).order_by(Evaluation.id.desc()))
        assert latest.input_snapshot['candidate']['skills'] == ['Rust']
        assert 'facts:2:project_experience' in latest.input_snapshot['evidence_catalog']
        assert 'HR 核实的新 Rust 项目' not in current.source.raw_resume


def test_s04_from_import_model_invite_mail_recovery_and_restart(env, monkeypatch, tmp_path):
    client, headers, factory = env
    record = import_sample(env, 4)
    with factory() as session: payload = _mock_evaluation(matching_input(session.get(Application, record.id)))
    for dimension in payload['dimensions']: dimension.update(score=0, references=[])
    for key, value in {'LLM_PROVIDER': 'openai_compatible', 'LLM_BASE_URL': 'https://model.example.test/v1', 'LLM_MODEL': 'synthetic', 'LLM_API_KEY': 'synthetic-only'}.items(): monkeypatch.setenv(key, value)
    monkeypatch.setattr('app.ai.client.httpx.post', lambda url, **kwargs: httpx.Response(401, request=httpx.Request('POST', url)))
    assert client.post(f'/api/applications/{record.id}/evaluate', headers=headers).status_code >= 400
    monkeypatch.setattr('app.ai.client.httpx.post', lambda url, **kwargs: httpx.Response(200, json={'choices': [{'message': {'content': json.dumps(payload)}}]}, request=httpx.Request('POST', url)))
    post(client, headers, f'/api/applications/{record.id}/evaluate')
    monkeypatch.setenv('LLM_PROVIDER', 'mock')
    enter_interview(client, headers, record.id)
    first = post(client, headers, f'/api/applications/{record.id}/schedule')
    with factory() as session:
        session.get(Interview, first['interview_id']).booking_token_expires_at = utc_now()-timedelta(seconds=1)
        session.commit()
    renewed = post(client, headers, f'/api/applications/{record.id}/schedule')
    assert renewed['interview_id'] == first['interview_id']
    assert client.post(first['booking_url']+'/confirm', data={'slot_id': first['slots'][0]['id']}).status_code == 404
    assert client.post(renewed['booking_url']+'/confirm', data={'slot_id': renewed['slots'][0]['id']}).status_code == 200
    post(client, headers, f"/api/interviews/{renewed['interview_id']}/feedback", next_action='offer', result='通过', comment='虚构面试核实')
    with factory() as session: offer_id = session.scalar(select(Offer.id).where(Offer.application_id == record.id))
    post(client, headers, f'/api/offers/{offer_id}/terms', version=0, company_name='虚构公司', job_title='虚构岗位', salary='虚构300元/天', subsidy='无',
         office_location='虚构办公区', report_location='虚构前台', report_time=(utc_now()+timedelta(days=7)).date().isoformat(), expires_at=(utc_now()+timedelta(days=3)).isoformat())
    post(client, headers, f'/api/offers/{offer_id}/approve', version=1)
    outcome = 'failed'; calls = []
    class SMTP:
        def __init__(self, *a, **k):
            if outcome == 'failed': raise ConnectionRefusedError('synthetic')
        def __enter__(self): return self
        def __exit__(self, *a): pass
        def starttls(self): pass
        def send_message(self, message):
            calls.append(message['To'])
            if outcome == 'unknown': raise TimeoutError('synthetic DATA timeout')
    monkeypatch.setenv('SMTP_HOST', 'smtp.example.test'); monkeypatch.setenv('SMTP_SSL', 'false')
    monkeypatch.setattr(offers.smtplib, 'SMTP', SMTP)
    for outcome in ['failed', 'unknown']:
        post(client, headers, f'/api/offers/{offer_id}/send')
        with factory() as session:
            assert session.get(Application, record.id).status == ApplicationStatus.OFFER_PENDING
            notification = session.scalar(select(Notification).where(Notification.channel == 'email').order_by(Notification.id.desc()))
            assert notification.status.value == outcome
    post(client, headers, f'/api/offers/{offer_id}/send'); assert len(calls) == 1
    post(client, headers, f'/api/notifications/{notification.id}/verify', delivered='false', note='HR 核对日志确认未送达')
    outcome = 'success'
    post(client, headers, f'/api/offers/{offer_id}/send'); assert len(calls) == 2
    with factory() as session: token = session.get(Offer, offer_id).candidate_token
    assert client.post(f'/offer/{token}/respond', data={'accepted': 'true'}).status_code == 200
    run_followups(factory)
    restart(factory, tmp_path)
    post(client, headers, f'/api/applications/{record.id}/onboarding', action='update', version=0, expected_date=utc_now().date().isoformat(), location='虚构前台')
    with factory() as session:
        assert session.get(Application, record.id).status == ApplicationStatus.ONBOARDING
        assert session.scalar(select(func.count(Evaluation.id))) == 2
        assert session.scalar(select(func.count(Notification.id)).where(Notification.channel == 'email')) == 3


def test_s06_replay_import_approval_reply_arrival_after_restart(env, tmp_path):
    client, headers, factory = env
    prepared = offer_from_normal_path(env, 6)
    post(client, headers, f'/api/offers/{prepared.id}/send')
    with factory() as session:
        offer = session.get(Offer, prepared.id); token = offer.candidate_token
        task = session.scalar(select(ImportTask)); import_id = task.id
    assert client.post(f'/offer/{token}/respond', data={'accepted': 'true'}).status_code == 200
    post(client, headers, f'/api/applications/{offer.application_id}/onboarding', action='update', version=0, expected_date=utc_now().date().isoformat(), location='虚构前台')
    with factory() as session: tasks = session.scalars(select(OnboardingTask)).all()
    for task in tasks: post(client, headers, f'/api/onboarding/{task.id}/complete')
    post(client, headers, f'/api/applications/{offer.application_id}/onboarding', action='join', version=1, actual_date=utc_now().date().isoformat(), note='HR 已现场核实')
    def counts():
        with factory() as session:
            return [session.scalar(select(func.count(model.id))) for model in [Application, CandidateSource, OfferApproval, Notification, OnboardingTask, AuditLog]]
    before = counts()
    for _ in range(2):
        assert post(client, headers, f'/api/import/tasks/{import_id}/confirm')['application_id'] == offer.application_id
        assert client.post(f'/api/offers/{offer.id}/approve', headers=headers, data={'version': 1}).status_code in {200, 409, 422}
        assert client.post(f'/offer/{token}/respond', data={'accepted': 'true'}).status_code == 200
        post(client, headers, f'/api/applications/{offer.application_id}/onboarding', action='join', version=1, actual_date=utc_now().date().isoformat(), note='HR 已现场核实')
        assert counts() == before
        restart(factory, tmp_path)


def test_booking_confirmation_rechecks_buffer_after_other_candidate_books(env):
    client, headers, _ = env
    bookings = []
    for sample in [1, 2]:
        record = import_sample(env, sample)
        enter_interview(client, headers, record.id)
        bookings.append(post(client, headers, f'/api/applications/{record.id}/schedule'))
    first_end = main.appointment_time(datetime.fromisoformat(bookings[0]['slots'][0]['end_at']))
    adjacent = next(s for s in bookings[1]['slots'] if main.appointment_time(datetime.fromisoformat(s['start_at'])) == first_end)
    available = next(s for s in bookings[1]['slots'] if main.appointment_time(datetime.fromisoformat(s['start_at'])) >= first_end+timedelta(minutes=15))
    assert client.post(bookings[0]['booking_url']+'/confirm', data={'slot_id': bookings[0]['slots'][0]['id']}).status_code == 200
    assert client.post(bookings[1]['booking_url']+'/confirm', data={'slot_id': adjacent['id']}).status_code == 409
    assert client.post(bookings[1]['booking_url']+'/confirm', data={'slot_id': available['id']}).status_code == 200


def test_s01_frozen_r01_two_exam_passes_to_independent_arrival(env):
    client, headers, factory = env
    offer = offer_from_normal_path(env, 1, passed_exams=True)
    post(client, headers, f'/api/offers/{offer.id}/send')
    with factory() as session: token = session.get(Offer, offer.id).candidate_token
    with TestClient(main.app) as candidate:
        assert candidate.post(f'/offer/{token}/respond', data={'accepted': 'true'}).status_code == 200
    endpoint = f'/api/applications/{offer.application_id}/onboarding'
    post(client, headers, endpoint, action='update', version=0, expected_date=utc_now().date().isoformat(), location='虚构前台')
    with factory() as session: tasks = session.scalars(select(OnboardingTask)).all()
    for task in tasks: post(client, headers, f'/api/onboarding/{task.id}/complete')
    with factory() as session: assert session.get(Application, offer.application_id).status == ApplicationStatus.ONBOARDING
    post(client, headers, endpoint, action='join', version=1, actual_date=utc_now().date().isoformat(), note='HR 现场核实')
    summary = client.get('/api/dashboard/summary').json()
    assert summary['funnel'] == {'applications': 1, 'screened': 1, 'assessments': 1, 'written': 1, 'interviews': 1, 'offers': 1, 'onboarded': 1}
